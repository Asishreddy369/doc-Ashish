import os
import tempfile
from django.test import TestCase, override_settings
import docx
import openpyxl
import fitz
from .extractor import extract_document_text, normalize_text

class TextExtractorTests(TestCase):
    def test_normalize_text(self):
        raw = "   Hello   World!   \r\n\r\n\r\n\r\nNext Line   "
        normalized = normalize_text(raw)
        self.assertEqual(normalized, "Hello World!\n\nNext Line")

    def test_docx_extraction(self):
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
            tmp_path = tmp.name

        try:
            doc = docx.Document()
            doc.add_paragraph("First paragraph content.")
            doc.add_paragraph("Second paragraph text.")
            doc.save(tmp_path)

            text, is_ocr, err, _blocks = extract_document_text(tmp_path, 'docx')
            self.assertIsNone(err)
            self.assertIn("First paragraph content.", text)
            self.assertIn("Second paragraph text.", text)
            self.assertFalse(is_ocr)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_xlsx_extraction(self):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp_path = tmp.name

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sales"
            ws['A1'] = "Product"
            ws['B1'] = "Price"
            ws['A2'] = "Laptop"
            ws['B2'] = 1200
            wb.save(tmp_path)

            text, is_ocr, err, _blocks = extract_document_text(tmp_path, 'xlsx')
            self.assertIsNone(err)
            self.assertIn("Product | Price", text)
            self.assertIn("Laptop | 1200", text)
            self.assertFalse(is_ocr)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_digital_pdf_extraction(self):
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
            tmp_path = tmp.name

        try:
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 50), "Document Management System Step 3 Text Extraction Test")
            doc.save(tmp_path)
            doc.close()

            text, is_ocr, err, _blocks = extract_document_text(tmp_path, 'pdf')
            self.assertIsNone(err)
            self.assertIn("Document Management System Step 3 Text Extraction Test", text)
            self.assertFalse(is_ocr)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    @override_settings(DOCVAULT_MAX_PDF_DIGITAL_PAGES=1, DOCVAULT_EXTRACTION_PROFILE="fast")
    def test_fast_pdf_extraction_limits_digital_pages(self):
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
            tmp_path = tmp.name

        try:
            pdf = fitz.open()
            first_page = pdf.new_page()
            first_page.insert_text(
                (50, 50),
                "First page GSTIN 29ABCDE1234F1Z5 with invoice and tax content for extraction",
            )
            second_page = pdf.new_page()
            second_page.insert_text((50, 50), "Second page should not be extracted in fast mode")
            pdf.save(tmp_path)
            pdf.close()

            text, is_ocr, err, _blocks = extract_document_text(tmp_path, 'pdf')
            self.assertIsNone(err)
            self.assertIn("First page GSTIN", text)
            self.assertNotIn("Second page should not be extracted", text)
            self.assertFalse(is_ocr)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    @override_settings(DOCVAULT_MAX_XLSX_ROWS=2, DOCVAULT_EXTRACTION_PROFILE="fast")
    def test_fast_xlsx_extraction_limits_rows(self):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            tmp_path = tmp.name

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Large"
            ws.append(["Header", "Value"])
            ws.append(["Included", "Row"])
            ws.append(["Skipped", "Row"])
            wb.save(tmp_path)

            text, is_ocr, err, _blocks = extract_document_text(tmp_path, 'xlsx')
            self.assertIsNone(err)
            self.assertIn("Header | Value", text)
            self.assertIn("Included | Row", text)
            self.assertIn("Sheet rows truncated", text)
            self.assertNotIn("Skipped | Row", text)
            self.assertFalse(is_ocr)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
