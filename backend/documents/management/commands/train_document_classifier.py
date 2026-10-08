import hashlib
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from documents.ml_classifier import (
    DEFAULT_MODEL_PATH,
    MultinomialTextClassifier,
    normalize_training_text,
)


SUPPORTED_EXTENSIONS = {
    ".pdf", ".jpg", ".jpeg", ".png", ".tiff", ".tif",
    ".docx", ".xlsx", ".txt", ".csv", ".log", ".bat",
}
DEFAULT_SKIP_LABELS = ["Uncertain", "Testing", "Yanvior Technolgies"]


def clean_filename_text(path):
    name = Path(path).stem
    name = re.sub(r"^[0-9a-fA-F-]{32,}_?", " ", name)
    name = re.sub(r"[_\-]+", " ", name)
    return name


def extract_fast_text(path):
    path = Path(path)
    ext = path.suffix.lower()
    parts = []

    if ext == ".pdf":
        import fitz

        with fitz.open(path) as doc:
            for page_index in range(min(len(doc), 2)):
                page = doc[page_index]
                text = page.get_text("text")
                if text:
                    parts.append(text)
    elif ext == ".docx":
        import docx

        document = docx.Document(path)
        parts.extend(p.text for p in document.paragraphs if p.text.strip())
        for table in document.tables[:5]:
            for row in table.rows[:30]:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    elif ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for sheet_name in wb.sheetnames[:4]:
            sheet = wb[sheet_name]
            for idx, row in enumerate(sheet.iter_rows(values_only=True)):
                if idx >= 80:
                    break
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    parts.append(" | ".join(values))
        wb.close()
    elif ext in {".txt", ".csv", ".log", ".bat"}:
        for encoding in ("utf-8", "cp1252", "latin-1"):
            try:
                parts.append(path.read_text(encoding=encoding, errors="ignore")[:12000])
                break
            except UnicodeError:
                continue
    else:
        return " ".join(parts)

    return normalize_training_text(" ".join(parts))


def extract_training_text(path, use_ocr=False):
    if use_ocr:
        from documents.extractor import extract_document_text

        text, _is_ocr, _err, _blocks = extract_document_text(str(path), Path(path).suffix.lower().lstrip("."))
        return normalize_training_text(text)
    return extract_fast_text(path)


def stable_sort_key(example):
    text, label = example
    return hashlib.sha256(f"{label}\0{text[:500]}".encode("utf-8", errors="ignore")).hexdigest()


def split_examples_for_evaluation(examples, eval_ratio=0.2):
    by_label = defaultdict(list)
    for example in examples:
        by_label[example[1]].append(example)

    train_examples = []
    eval_examples = []
    for label, label_examples in by_label.items():
        ordered = sorted(label_examples, key=stable_sort_key)
        if len(ordered) < 5 or eval_ratio <= 0:
            train_examples.extend(ordered)
            continue
        eval_count = max(1, int(round(len(ordered) * eval_ratio)))
        eval_examples.extend(ordered[:eval_count])
        train_examples.extend(ordered[eval_count:])

    if len({label for _text, label in train_examples}) < 2:
        return examples, []
    return train_examples, eval_examples


def evaluate_model(train_examples, eval_examples):
    if not eval_examples:
        return {}

    model = MultinomialTextClassifier()
    model.fit(train_examples)

    total = 0
    correct = 0
    confusion = Counter()
    per_label = defaultdict(lambda: {"total": 0, "correct": 0})

    for text, expected_label in eval_examples:
        prediction = model.predict(text)
        predicted_label = prediction["category"] if prediction else "Unclassified"
        total += 1
        per_label[expected_label]["total"] += 1
        if predicted_label == expected_label:
            correct += 1
            per_label[expected_label]["correct"] += 1
        else:
            confusion[(expected_label, predicted_label)] += 1

    return {
        "strategy": "deterministic_label_holdout",
        "eval_documents": total,
        "accuracy": round(correct / total, 4) if total else None,
        "per_label": {
            label: {
                "total": values["total"],
                "correct": values["correct"],
                "accuracy": round(values["correct"] / values["total"], 4) if values["total"] else None,
            }
            for label, values in sorted(per_label.items())
        },
        "top_confusions": [
            {"expected": expected, "predicted": predicted, "count": count}
            for (expected, predicted), count in confusion.most_common(10)
        ],
    }


class Command(BaseCommand):
    help = "Train the local lightweight ML document classifier from a folder where child folders are labels."

    def add_arguments(self, parser):
        parser.add_argument(
            "dataset_root",
            nargs="?",
            default=r"C:\Users\HP\Downloads\Yanvior Technology",
            help="Folder containing one subfolder per document category.",
        )
        parser.add_argument(
            "--from-db",
            action="store_true",
            help="Also train from existing database documents that have extracted text and a non-generic category.",
        )
        parser.add_argument(
            "--model-path",
            default=str(DEFAULT_MODEL_PATH),
            help="Where to write the trained JSON model.",
        )
        parser.add_argument(
            "--use-ocr",
            action="store_true",
            help="Use the full OCR extractor for image/scanned files. Slower, but can improve scanned-doc training.",
        )
        parser.add_argument(
            "--limit-per-label",
            type=int,
            default=0,
            help="Optional maximum number of documents to train from each label.",
        )
        parser.add_argument(
            "--min-docs-per-label",
            type=int,
            default=2,
            help="Skip labels with fewer successfully extracted documents.",
        )
        parser.add_argument(
            "--skip-label",
            action="append",
            default=DEFAULT_SKIP_LABELS.copy(),
            help="Folder label to skip. Can be repeated.",
        )
        parser.add_argument(
            "--eval-ratio",
            type=float,
            default=0.2,
            help="Deterministic holdout ratio per label for training evaluation metadata.",
        )

    def handle(self, *args, **options):
        dataset_root = Path(options["dataset_root"]).expanduser()
        if not dataset_root.exists():
            raise CommandError(f"Dataset folder not found: {dataset_root}")

        skip_labels = {label.lower() for label in options["skip_label"]}
        examples = []
        counts = {}
        failed = []

        label_dirs = [
            child for child in sorted(dataset_root.iterdir())
            if child.is_dir() and child.name.lower() not in skip_labels
        ]

        if len(label_dirs) < 2:
            raise CommandError("Need at least two category folders to train.")

        for label_dir in label_dirs:
            label = label_dir.name
            trained_for_label = 0
            files = [
                path for path in label_dir.rglob("*")
                if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
            ]

            for path in files:
                if options["limit_per_label"] and trained_for_label >= options["limit_per_label"]:
                    break
                try:
                    text = extract_training_text(path, use_ocr=options["use_ocr"])
                    if len(text) < 20:
                        continue
                    examples.append((text, label))
                    trained_for_label += 1
                except Exception as exc:
                    failed.append({"file": str(path), "error": str(exc)})

            counts[label] = trained_for_label
            self.stdout.write(f"{label}: {trained_for_label} training document(s)")

        db_count = 0
        if options["from_db"]:
            from documents.models import Document

            generic_labels = {"", "Unclassified", "Other", "Other Documents"}
            queryset = Document.objects.exclude(category__in=generic_labels).exclude(extracted_text__isnull=True).exclude(extracted_text="")
            for doc in queryset.iterator():
                text = normalize_training_text(" ".join([
                    doc.extracted_text or "",
                ]))
                if len(text) < 20:
                    continue
                examples.append((text, doc.category))
                counts[doc.category] = counts.get(doc.category, 0) + 1
                db_count += 1
            self.stdout.write(f"Database: {db_count} training document(s) with extracted text")

        min_docs = options["min_docs_per_label"]
        examples = [
            (text, label)
            for text, label in examples
            if counts.get(label, 0) >= min_docs
        ]
        usable_labels = sorted({label for _text, label in examples})

        if len(usable_labels) < 2:
            raise CommandError("Training stopped: fewer than two labels had enough extractable documents.")

        train_examples, eval_examples = split_examples_for_evaluation(examples, options["eval_ratio"])
        evaluation = evaluate_model(train_examples, eval_examples)

        model = MultinomialTextClassifier()
        model.fit(examples)
        model.trained_at = timezone.now().isoformat()
        model.training_summary = {
            "dataset_root": str(dataset_root),
            "labels": usable_labels,
            "documents_per_label": {label: counts[label] for label in usable_labels},
            "total_documents": len(examples),
            "failed_documents": len(failed),
            "used_ocr": bool(options["use_ocr"]),
            "database_documents": db_count,
            "evaluation": evaluation,
        }
        model.save(options["model_path"])

        self.stdout.write(self.style.SUCCESS(f"Saved ML classifier: {options['model_path']}"))
        self.stdout.write(self.style.SUCCESS(f"Trained labels: {', '.join(usable_labels)}"))
        self.stdout.write(self.style.SUCCESS(f"Training documents used: {len(examples)}"))
        if evaluation:
            self.stdout.write(self.style.SUCCESS(f"Holdout accuracy: {evaluation['accuracy']} on {evaluation['eval_documents']} document(s)"))
        if failed:
            self.stdout.write(self.style.WARNING(f"Skipped {len(failed)} file(s) due to extraction errors."))
